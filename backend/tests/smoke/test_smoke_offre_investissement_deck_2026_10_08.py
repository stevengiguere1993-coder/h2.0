"""Smoke — pitch deck (offre d'investissement .pptx) v3, Phil 2026-10-08 :
« le résultat qui sort demande beaucoup d'ajustement […] le but est que
quand je génère, je n'ai pas ou presque rien à changer ensuite ».

- l'assistant reçoit des intrants proposés + un aperçu des chiffres ;
- la génération remplit chaque cible du gabarit par son nom (0 oubli),
  sans fuite des chiffres du deal qui a servi de gabarit (2420 Pie-IX) ;
- l'échéancier est calculé (jalons datés, losanges placés dans l'ordre) ;
- prêteur B (avec balance de vente) et institution traditionnelle.
"""
from __future__ import annotations

import base64
import io
import json

from app.models.lead_analysis import LeadAnalysis
from app.services import offre_investissement_deck as deck

from tests.smoke.conftest import TestSessionLocal

_FUITES = (
    "Pie-IX", "Hochelaga", "1 800 000$", "1 000 984", "316  765", "2 834 515",
    "3 334 723", "Septembre-Octobre 2026", "Balance de vente de 500 000$",
    "Août 2027", "Décembre 2028", "7,85%", "+200 000$", "1,8 M$",
)


def _png() -> bytes:
    from PIL import Image

    img = Image.new("RGB", (900, 600), (30, 90, 160))
    b = io.BytesIO()
    img.save(b, "PNG")
    return b.getvalue()


def _creer(run, **kw) -> int:
    base = dict(
        address="4321 rue des Érables", city="Longueuil", asking_price=1_250_000, nb_logements=6,
        typology_json=json.dumps({"3.5": 2, "4.5": 4}), revenus_bruts=78_000, taxes_municipales=9_800,
        taxes_scolaires=900, assurances=5_200, energie=0, depenses_autres=0,
        loyers_projetes_json=json.dumps({"3.5": 1_250, "4.5": 1_550}), taux_interet_refi_pct=4.0,
        tga_pct=4.0, duree_projet_annees=2, travaux_estimes=180_000, annee_construction=1965,
        nb_stationnements=4, tri_capital_injecte=250_000.0, tri_pct_investisseur=0.5,
    )
    base.update(kw)

    async def _c():
        async with TestSessionLocal() as s:
            rec = LeadAnalysis(**base)
            s.add(rec)
            await s.commit()
            await s.refresh(rec)
            return rec.id

    return run(_c())


def _textes(prs) -> str:
    morceaux = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                morceaux.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        morceaux.append(c.text_frame.text)
            if getattr(sh, "has_chart", False) and sh.has_chart and sh.chart.has_title:
                morceaux.append(sh.chart.chart_title.text_frame.text)
    return "\n".join(morceaux)


def _shape(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def test_formats():
    assert deck.money(1_800_000) == "1 800 000$"
    assert deck.money(-534_884.4) == "-534 884$"
    assert deck.money(210_000, signe=True) == "+210 000$"
    assert deck.money_M(1_800_000) == "1,8 M$" and deck.money_M(2_000_000) == "2 M$"
    assert deck.money_M(950_000) == "950 000$"
    assert deck.money_an(130_824) == "130 824 $/an"
    assert deck.pct(0.0785, 2) == "7,85%" and deck.pct(0.8) == "80%"
    assert deck.pct(0.375, 1, espace=True) == "37,5 %"
    assert deck.ratio(1.1) == "1,10"
    assert deck._typologie_texte(json.dumps({"6.5": 6, "5.5": 2})) == "2 x 5 ½ + 6 x 6 ½"


def test_defaults_puis_generation_preteur_b(client, auth_headers, run):
    fid = _creer(run, strategie_acquisition="preteur_b", balance_vente_montant=150_000, balance_vente_taux_pct=5.0)
    base = f"/api/v1/lead-analyses/{fid}"
    # Sans analyse : l'assistant le dit, la génération refuse proprement.
    d0 = client.get(f"{base}/offre-investissement/defaults", headers=auth_headers).json()
    assert d0["analysis_ready"] is False and d0["inputs"] is None and "analyse" in d0["raison"].lower()
    assert client.post(f"{base}/offre-investissement", headers=auth_headers, json={"inputs": {}}).status_code == 409

    assert client.post(f"{base}/run-financial-analysis", headers=auth_headers).status_code == 200
    r = client.get(f"{base}/offre-investissement/defaults", headers=auth_headers)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["analysis_ready"] is True and d["template_version"] == "horizon_v3"
    inp = d["inputs"]
    assert "Longueuil" in inp["tagline"] and len(inp["bullets"]) == 4 and len(inp["leviers"]) == 4
    assert set(inp["jalons"]) == {"m1_1", "m1_2", "m2_1", "m2_2", "m2_3", "m2_4", "m3_1"}
    assert inp["investissement_requis"] == 250_000 and inp["pct_parts"] == 0.5
    assert inp["unites_note"] == "4 stationnements" and inp["frais_energetiques_text"] == "Locataires"
    ap = d["apercu"]
    assert ap["prix"] == 1_250_000 and ap["nb_logements"] == 6 and ap["balance_vente"] == 150_000
    assert [h["annee"] for h in ap["tri"]] == [2, 7, 12] and ap["capital"] == 250_000
    assert [s["key"] for s in d["photos"]["slots"]][:2] == ["cover", "resume_1"]
    assert "Toit" in d["renovations_catalogue"]
    # Aperçu recalculé avec un autre capital / % de parts.
    d2 = client.get(f"{base}/offre-investissement/defaults?investissement_requis=400000&pct_parts=0.4", headers=auth_headers).json()
    assert d2["apercu"]["capital"] == 400_000 and d2["apercu"]["pct_parts"] == 0.4
    assert d2["inputs"]["investissement_requis"] == 400_000

    # Génération avec quelques ajustements et une photo de couverture.
    inp["nb_etages"] = 3
    inp["renovations"] = ["Toit", "Chauffe-eau", "Cuisine complète"]
    inp["comparable_url"] = "https://www.centris.ca/fr/exemple"
    inp["jalons"]["m2_1"]["label"] = "Ententes signées"
    r = client.post(
        f"{base}/offre-investissement", headers=auth_headers,
        json={"inputs": inp, "photos": {"cover": {"base64_data": base64.b64encode(_png()).decode()}}},
    )
    assert r.status_code == 200, r.text
    assert r.headers["X-Deck-Misses"] == "0" and r.headers["X-Template-Version"] == "horizon_v3"
    assert "presentationml" in r.headers["content-type"]

    from pptx import Presentation

    prs = Presentation(io.BytesIO(r.content))
    assert len(prs.slides) == 16
    texte = _textes(prs)
    for fuite in _FUITES:
        assert fuite not in texte, f"fuite du gabarit : {fuite!r}"
    assert "4321 rue des Érables, Longueuil" in texte
    assert "1 250 000$" in texte and "Immeuble de 6 logements" in texte
    assert "2 x 3 ½ + 4 x 4 ½" in texte and "Ententes signées" in texte
    assert "Prêteur B" in texte and "Balance de vente de 150 000$" in texte
    assert "Toit" in texte and "Cuisine complète" in texte
    assert "TENDANCES — LONGUEUIL" in texte
    assert "(50% des parts)" in texte and "50% des parts du projet" in texte
    # TRI imprimé = TRI de l'aperçu (mêmes intrants).
    tri2 = ap["tri"][0]["tri"]
    assert deck.pct(tri2, 1) in texte and deck.money(ap["tri"][0]["cash"]) in texte
    # Diapo 6 : losanges dans l'ordre chronologique, dans le tableau.
    s6 = prs.slides[5]
    table = _shape(s6, "Table 1")
    xs = [_shape(s6, f"Flowchart: Decision {n}").left for n in (34, 35, 36, 37, 38, 39, 40)]
    assert xs == sorted(xs), xs
    assert table.left <= xs[0] and xs[-1] <= table.left + table.width
    for n in (34, 35, 36, 37, 38, 39, 40):
        sh = _shape(s6, f"Flowchart: Decision {n}")
        assert table.top <= sh.top <= table.top + table.height
    # Repère « aujourd'hui » dans la première colonne de mois.
    today = _shape(s6, "Straight Connector 74")
    assert table.left <= today.left <= table.left + table.width
    # Photo de couverture remplacée (nouvelle image), pastille ailleurs.
    cover = _shape(prs.slides[0], "Picture 33")
    assert cover.image.content_type in ("image/jpeg", "image/png") and len(cover.image.blob) < 200_000
    pastille = _shape(prs.slides[0], "Picture 33").image.blob
    gabarit = _shape(Presentation(str(deck.template_path())).slides[0], "Picture 33").image.blob
    assert pastille != gabarit


def test_generation_institution_traditionnelle(client, auth_headers, run):
    fid = _creer(
        run, address="77 avenue du Parc", city="Sherbrooke", strategie_acquisition="traditionnel",
        projection_horizon_annees=5, programme_achat="aph_100",
    )
    base = f"/api/v1/lead-analyses/{fid}"
    assert client.post(f"{base}/run-financial-analysis", headers=auth_headers).status_code == 200
    d = client.get(f"{base}/offre-investissement/defaults", headers=auth_headers).json()
    assert d["analysis_ready"] is True and d["apercu"]["strategie"] == "Institution financière"
    assert [h["annee"] for h in d["apercu"]["tri"]] == [5, 10, 15]
    r = client.post(f"{base}/offre-investissement", headers=auth_headers, json={"inputs": d["inputs"]})
    assert r.status_code == 200, r.text
    assert r.headers["X-Deck-Misses"] == "0"

    from pptx import Presentation

    prs = Presentation(io.BytesIO(r.content))
    texte = _textes(prs)
    for fuite in _FUITES:
        assert fuite not in texte, f"fuite du gabarit : {fuite!r}"
    assert "Financement institutionnel" in texte and "Institution financière" in texte
    assert "Refinancement 5 ans" in texte and "Année 15" in texte
    # Pas de balance de vente → l'encart est retiré ; pas de gain comparable → callout retiré.
    noms8 = {s.name for s in prs.slides[7].shapes}
    assert "Rectangle 1" not in noms8
    assert "Amorti" in texte


def test_photos_par_piece_jointe_et_erreurs(client, auth_headers, run):
    from app.models.lead_analysis import LeadAnalysisAttachment

    fid = _creer(run, address="9 rue des Photos", city="Laval")
    base = f"/api/v1/lead-analyses/{fid}"
    assert client.post(f"{base}/run-financial-analysis", headers=auth_headers).status_code == 200

    async def _att():
        async with TestSessionLocal() as s:
            a = LeadAnalysisAttachment(lead_analysis_id=fid, filename="facade.png", content_type="image/png",
                                       size_bytes=10, blob=_png())
            s.add(a)
            await s.commit()
            await s.refresh(a)
            return a.id

    aid = run(_att())
    d = client.get(f"{base}/offre-investissement/defaults", headers=auth_headers).json()
    assert [a["id"] for a in d["photos"]["attachments"]] == [aid]
    r = client.post(f"{base}/offre-investissement", headers=auth_headers,
                    json={"inputs": {}, "photos": {"avant": {"attachment_id": aid}, "apres": {"attachment_id": aid}}})
    assert r.status_code == 200 and r.headers["X-Deck-Misses"] == "0"
    # Pièce jointe d'une autre fiche / emplacement inconnu → 400.
    assert client.post(f"{base}/offre-investissement", headers=auth_headers,
                       json={"inputs": {}, "photos": {"avant": {"attachment_id": aid + 999}}}).status_code == 400
    assert client.post(f"{base}/offre-investissement", headers=auth_headers,
                       json={"inputs": {}, "photos": {"inconnu": {"attachment_id": aid}}}).status_code == 400
